"""Build the Project Final deck (1 Nov video round) in the organizer template's look.

Run: uv run --with python-pptx python proposal/project_final/build_pptx.py [--script PATH]

Geometry, colours and fonts are measured from the organizer's
"Template Project Final Healthkathon 2026" PDF (959.76 x 540 pt pages).
Slide copy follows Video Script v1; narration and stage directions are read from
the script file and written into the speaker notes, so a script edit only needs a rebuild.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "out" / "TEPAT Project Final.pptx"
DEFAULT_SCRIPT = Path("C:/Users/Bani/Documents/THE HIVE/07-Healthkathon/09-Pitch/Video Script v1.md")

NAVY = "172540"
NAVY2 = "22304f"
GREEN = "22a05a"
MINT = "2ed3a0"
PALE = "d6e4dd"
PANEL = "f2f6f4"
PANEL_LINE = "c9d6cf"
TINT = "e3f5ea"
TINT_LINE = "bfe3cd"
GRAY = "7c8a96"
PH = "8a97a3"
BODY = "2b3a4a"
WHITE = "ffffff"

HEAD = "Cambria"
SANS = "Calibri"
TAGLINE = "Detect Smarter, Protect JKN - Smarter Protection for a Stronger JKN"
DEMO_URL = "https://tepat-jkp.streamlit.app"
REPO_URL = "https://github.com/rabbaniputrowi-alt/TEPAT"

BANNED = {"\u2014": "em dash", "\u2013": "en dash", "\u00b7": "middle dot", "\u2026": "ellipsis"}


# ----------------------------------------------------------------- primitives

def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.upper())


def box(slide, x, y, w, h, fill=None, line=None, radius=0.0, dash=False, lw=1.0, alpha=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Pt(x), Pt(y), Pt(w), Pt(h))
    if radius:
        shape.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(alpha * 1000)))
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(lw)
        if dash:
            shape.line.dash_style = MSO_LINE.DASH
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.text = ""
    return shape


def circle(slide, x, y, d, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y), Pt(d), Pt(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _style_run(run, size, bold=False, italic=False, color=BODY, font=SANS, spacing=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", font)
    if spacing:
        rpr.set("spc", str(int(spacing * 100)))


def text(slide_or_shape, x=None, y=None, w=None, h=None, paras=(), anchor="top", inset=0.0, autofit=False):
    """paras: list of dicts. Keys: t (str) or runs [(str, overrides)], size, bold, italic,
    color, font, align, after, before, spacing, bullet, line, link."""
    if x is None:  # write into an existing shape
        shape = slide_or_shape
    else:
        shape = slide_or_shape.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Pt(inset))
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[p.get("align", "l")]
        if "after" in p:
            para.space_after = Pt(p["after"])
        if "before" in p:
            para.space_before = Pt(p["before"])
        para.line_spacing = p.get("line", 1.05)
        base = {k: p[k] for k in ("size", "bold", "italic", "color", "font", "spacing") if k in p}
        if p.get("link"):
            # Link on the text box itself, so PowerPoint keeps the run colour instead of the theme's blue.
            shape.click_action.hyperlink.address = p["link"]
        base.setdefault("size", 13)
        runs = p.get("runs") or [(p.get("t", ""), {})]
        for s, over in runs:
            r = para.add_run()
            r.text = s
            _style_run(r, **{**base, **over})
        if p.get("bullet"):
            ppr = para._p.get_or_add_pPr()
            ind = p.get("indent", 13)
            ppr.set("marL", str(int(Pt(ind))))
            ppr.set("indent", str(int(-Pt(ind))))
            for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
                for el in ppr.findall(qn(tag)):
                    ppr.remove(el)
            buclr = etree.SubElement(ppr, qn("a:buClr"))
            etree.SubElement(buclr, qn("a:srgbClr")).set("val", p.get("bullet_color", GREEN).upper())
            etree.SubElement(ppr, qn("a:buSzPct")).set("val", "100000")
            etree.SubElement(ppr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(ppr, qn("a:buChar")).set("char", "\u2022")
    if autofit:
        bodypr = tf._txBody.find(qn("a:bodyPr"))
        etree.SubElement(bodypr, qn("a:normAutofit"))
    return shape


def label(slide, x, y, s, color=GREEN, w=420, size=12.2):
    text(slide, x, y, w, 18, [{"t": s.upper(), "size": size, "bold": True, "color": color, "spacing": 2.4}])


def content_slide(prs, badge, title, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(WHITE)
    b = box(s, 43.6, 36.4, 72.0, 57.6, fill=GREEN, radius=10)
    text(b, paras=[{"t": badge, "size": 33.9, "bold": True, "color": WHITE, "font": HEAD, "align": "c"}], anchor="middle")
    text(s, 140.5, 38.0, 780, 50, [{"t": title, "size": 31.7, "bold": True, "color": NAVY, "font": HEAD}], anchor="middle")
    text(s, 43.2, 510.5, 200, 16, [{"t": "Project Final Healthkathon 2026", "size": 10.1, "bold": True, "color": GREEN}])
    text(s, 237.7, 510.5, 520, 16, [{"t": TAGLINE, "size": 10.1, "italic": True, "color": GRAY}])
    text(s, 856, 510.5, 60, 16, [{"t": str(page), "size": 10.1, "color": GRAY, "align": "r"}])
    return s


def navy_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY)
    s.shapes.add_picture(str(ASSETS / "healthkathon_logo.png"), Pt(734.4), Pt(36.0), Pt(187.2), Pt(64.1))
    return s


def navy_card(s, x, y, w, h, head, body_paras):
    box(s, x, y, w, h, fill=GREEN, line=MINT, radius=7, alpha=45.098)
    label(s, x + 17.6, y + 12.5, head, color=MINT, w=w - 30, size=12.2)
    text(s, x + 17.6, y + 38, w - 33, h - 44, body_paras)


def notes(slide, s: str):
    slide.notes_slide.notes_text_frame.text = s


# ----------------------------------------------------------------- script notes

def read_script_notes(path: Path) -> dict[int, str]:
    """Section number -> notes text: header, stage direction, narration (verbatim)."""
    if not path.exists():
        return {}
    out: dict[int, str] = {}
    sec = None
    buf: list[str] = []
    head = ""

    def flush():
        if sec is not None:
            out[sec] = head + "\n\n" + "\n\n".join(x for x in buf if x)

    speaker = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^## (\d+)\. (.+?) \((\d+:\d+) to (\d+:\d+)\)", raw)
        if m:
            flush()
            sec = int(m.group(1))
            buf = []
            head = f"Bagian {sec}. {m.group(2).split(':')[0]}. Waktu {m.group(3)} s.d. {m.group(4)}."
            continue
        if raw.startswith("## "):
            flush()
            sec = None
            continue
        if sec is None:
            continue
        line = raw.strip()
        m2 = re.match(r"^\*\*Narasi \((\w+)\)\*\*", line)
        if m2:
            buf.append(f"NARASI ({m2.group(1)}):")
            continue
        m3 = re.match(r"^\*\*Segment (\w): (.+?)\*\*$", line)
        if m3:
            buf.append(f"SEGMEN {m3.group(1)}: {m3.group(2)}")
            continue
        if line.startswith(">") and not line.startswith("> [!"):
            t = line.lstrip("> ").strip()
            if t:
                buf.append(t)
            continue
        if line.startswith("*") and not line.startswith("**") and line.endswith("*") and not buf:
            buf.append("ARAHAN: " + line.strip("*"))
    flush()
    return out


# ----------------------------------------------------------------- slides

def build(script: Path) -> Presentation:
    N = read_script_notes(script)
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)

    # Cover --------------------------------------------------------------
    s = navy_slide(prs)
    text(s, 64.8, 116, 600, 26, [{"t": "PROJECT FINAL HEALTHKATHON 2026", "size": 20.2, "bold": True, "color": MINT}])
    text(s, 64.8, 150, 820, 80, [{"t": "TEPAT", "size": 56.2, "bold": True, "color": WHITE, "font": HEAD}])
    text(s, 64.8, 240, 820, 60, [
        {"t": "Pemeriksaan Kepatuhan Tepat Sasaran", "size": 22, "color": WHITE, "after": 4},
        {"t": TAGLINE + ": Efisiensi Risiko Program JKN", "size": 15, "italic": True, "color": PALE},
    ])
    navy_card(s, 65.2, 382.1, 198.0, 86.4, "Nama solusi", [{"t": "TEPAT", "size": 17.3, "bold": True, "color": WHITE}])
    navy_card(s, 277.6, 382.1, 198.0, 86.4, "Nama tim", [
        {"t": "Johannes Kepler", "size": 17.3, "bold": True, "color": WHITE},
        {"t": "SMA Negeri 34 Jakarta", "size": 12.5, "color": PALE}])
    navy_card(s, 490.0, 382.1, 260.0, 86.4, "Kategori", [
        {"t": "Efisiensi Risiko pada Pemberi Kerja", "size": 15.5, "bold": True, "color": WHITE}])
    text(s, 64.8, 490, 820, 20, [{"t": "Seluruh data di video ini sintetis. Tidak ada data peserta JKN riil.",
                                  "size": 12.5, "italic": True, "color": PALE}])
    notes(s, N.get(1, ""))

    # 1 Perkenalan Tim ----------------------------------------------------
    s = content_slide(prs, "1", "Perkenalan Tim", 1)
    text(s, 43.2, 110, 880, 24, [{"runs": [("Nama tim: ", {}), ("Johannes Kepler, SMA Negeri 34 Jakarta", {"bold": False}),
                                            ("      |      Kategori: ", {}), ("Efisiensi Risiko pada Pemberi Kerja", {"bold": False})],
                                  "size": 15.9, "bold": True, "color": NAVY}])
    members = [
        ("Rabbani Putrowi", "Ketua: data, model, dan aplikasi",
         "Membangun data sintetis, model, benchmark, dan aplikasi TEPAT."),
        ("Raffa Maliq Revaya", "Regulasi dan kepatuhan",
         "Bertanggung jawab atas riset regulasi kepatuhan pemberi kerja dan dasar hukum setiap alasan."),
        ("Ahmad Fahri Fauzi", "Produk dan pengalaman pengguna",
         "Bertanggung jawab atas pengalaman pemeriksa di aplikasi dan penyampaian solusi."),
    ]
    for i, (name, role, resp) in enumerate(members):
        x0 = 43.6 + i * 295.2
        box(s, x0, 143, 280.8, 290, fill=PANEL, line=PANEL_LINE, radius=7)
        ph = box(s, x0 + 21.2, 160, 112, 136, fill=WHITE, line=PANEL_LINE, radius=7, dash=True, lw=1.25)
        text(ph, paras=[{"t": f"Foto\n{name.split()[0]}", "size": 13, "italic": True, "color": PH, "align": "c"}], anchor="middle")
        label(s, x0 + 146, 162, f"Anggota {i + 1}", w=130, size=10.8)
        text(s, x0 + 146, 182, 124, 112, [
            {"t": name, "size": 16, "bold": True, "color": NAVY, "after": 6, "line": 0.95},
            {"t": role, "size": 12.5, "color": BODY, "line": 1.0}])
        rb = box(s, x0 + 21.2, 308, 238.4, 112, fill=WHITE, line=PANEL_LINE, radius=7)
        text(s, x0 + 32, 318, 217, 94, [{"t": resp, "size": 15, "color": BODY, "line": 1.08}])
    box(s, 43.6, 444, 871.2, 50, fill=TINT, line=TINT_LINE, radius=7)
    label(s, 61.2, 462, "Pengalaman tim", w=150, size=11)
    for j, cred in enumerate(["DATATHON 2026 - Finalist Ristek UI", "Algoverse Researcher", "JISF Silver Award"]):
        text(s, 210 + j * 238, 457, 236, 24, [{"t": cred, "size": 14.5, "bold": True, "color": NAVY}])
    notes(s, N.get(2, ""))

    # 2 Judul -------------------------------------------------------------
    s = content_slide(prs, "2", "Judul Ide, Deskripsi, dan Tujuan", 2)
    box(s, 43.6, 119.2, 388.8, 374.5, fill=NAVY, radius=7)
    label(s, 68.4, 138, "Nama solusi", color=MINT)
    text(s, 68.4, 158, 350, 70, [
        {"t": "TEPAT", "size": 30.3, "bold": True, "color": WHITE, "font": HEAD},
        {"t": "Pemeriksaan Kepatuhan Tepat Sasaran", "size": 15, "color": PALE}])
    label(s, 68.4, 240, "Satu kalimat nilai utama", color=MINT)
    text(s, 68.4, 262, 345, 90, [{"t": "TEPAT membantu tim kepatuhan BPJS Kesehatan memilih pemberi kerja yang "
                                       "diperiksa lebih dulu, agar jam pemeriksa yang terbatas dipakai untuk kasus "
                                       "dengan bukti terkuat.", "size": 15, "italic": True, "color": PALE, "line": 1.08}])
    label(s, 68.4, 364, "Kategori dan sub kategori", color=MINT)
    text(s, 68.4, 386, 345, 100, [
        {"t": "Efisiensi Risiko pada Pemberi Kerja", "size": 15, "bold": True, "color": WHITE, "after": 5},
        {"runs": [("Inti: ", {"bold": True}), ("1 Pendaftaran pekerja sebagian, 2 Under-reporting upah, "
                                                "3 Penundaan setoran iuran, 5 Manipulasi data mutasi.", {})],
         "size": 12.2, "color": WHITE, "after": 3, "line": 1.05},
        {"runs": [("Pendukung: ", {"bold": True}), ("4 Misklasifikasi status kerja.", {})], "size": 12.2, "color": WHITE}])
    label(s, 453.8, 119, "Deskripsi singkat")
    box(s, 453.6, 140, 463.0, 165, fill=WHITE, line=PANEL_LINE, radius=7)
    text(s, 468, 152, 436, 145, [
        {"t": "Mencocokkan data kepesertaan dan iuran BPJS Kesehatan dengan file pertukaran BPJS Ketenagakerjaan, "
              "lalu menyusun daftar tindak lanjut sesuai jam pemeriksa yang tersedia.", "size": 16.5, "color": BODY,
         "after": 10, "line": 1.08},
        {"t": "Setiap peringkat disertai alasan yang bisa diperiksa dan dasar pasalnya.", "size": 16.5, "color": BODY,
         "line": 1.08}])
    label(s, 453.8, 316, "Tujuan")
    box(s, 453.6, 337, 463.0, 156.7, fill=WHITE, line=PANEL_LINE, radius=7)
    goals = ["Lebih banyak temuan per jam pemeriksa.",
             "Lebih sedikit pemanggilan pemberi kerja yang patuh.",
             "Keputusan tetap di tangan pemeriksa."]
    for k, g in enumerate(goals):
        yy = 354 + k * 44
        c = circle(s, 470, yy, 26, GREEN)
        text(c, paras=[{"t": str(k + 1), "size": 13, "bold": True, "color": WHITE, "align": "c"}], anchor="middle")
        text(s, 508, yy - 1, 395, 28, [{"t": g, "size": 15, "color": BODY}], anchor="middle")
    notes(s, N.get(3, ""))

    # 3 Latar dan dampak ---------------------------------------------------
    s = content_slide(prs, "3", "Latar Belakang dan Dampak", 3)
    label(s, 43.2, 114, "Latar belakang")
    box(s, 43.2, 135, 424.8, 340, fill=WHITE, line=PANEL_LINE, radius=7)
    bl = dict(size=14.2, color=BODY, bullet=True, after=8, line=1.05)
    text(s, 58, 147, 396, 320, [
        {"t": "Iuran pemberi kerja yang tidak masuk ikut menekan dana JKN, dan pemeriksa harus memilih badan usaha "
              "mana yang diperiksa lebih dulu.", "size": 15, "bold": True, "color": NAVY, "after": 9, "line": 1.05},
        {"t": "Iuran terkumpul Rp118 triliun, klaim Rp132 triliun.", **bl},
        {"t": "Lebih dari 320 ribu badan usaha sektor formal, 128 kantor cabang.", **bl},
        {"t": "7 Sep 2026: PKS pengawasan bersama BPJS Kesehatan dan BPJS Ketenagakerjaan. Langkah berikut yang "
              "dinyatakan: memakai data pertukaran untuk menentukan sasaran pemeriksaan. 2.429 badan usaha sudah "
              "diperiksa bersama sampai Juli 2026.", **bl},
        {"t": "Cara memilih sasaran saat ini tidak dipublikasikan.", **bl},
        {"t": "Terdampak: pekerja (kepesertaan dihentikan sementara bila iuran tidak dibayar, Perpres 59/2024 Ps. 42), "
              "dana JKN, dan pemberi kerja yang patuh.", **bl},
    ])
    label(s, 489.9, 114, "Dampak (data sintetis)")
    box(s, 489.6, 135, 427.0, 178, fill=WHITE, line=PANEL_LINE, radius=7)
    text(s, 504, 146, 400, 160, [
        {"t": "Cabang uji B, 5 seed, anggaran jam pemeriksaan yang sama. Pembanding: aturan yang kami susun sendiri, "
              "bukan cara BPJS saat ini.", "size": 13.6, "color": BODY, "after": 6, "line": 1.05},
        {"t": "Dengan jalur acak 8%, presisi 84,8%. Angka rupiah adalah potensi sintetis, bukan uang yang tertagih.",
         "size": 13.6, "color": BODY, "after": 6, "line": 1.05},
        {"runs": [("Kondisi gagal: ", {"bold": True, "color": NAVY}),
                  ("bila pelanggaran jarang (prevalensi 5%), presisi turun ke 34,3%.", {})],
         "size": 13.6, "color": BODY, "line": 1.05},
    ])
    tiles = [("95,3%", "Presisi TEPAT tanpa jalur acak", "Aturan pembanding: 35,4%"),
             ("33,7%", "Recall TEPAT", "Aturan pembanding: 8,8%"),
             ("2,27x", "Potensi Rp per jam", "Rp8,86 jt vs Rp3,91 jt")]
    for k, (num, cap, base) in enumerate(tiles):
        x0 = 490.0 + k * 144.0
        box(s, x0, 323, 139.0, 152, fill=TINT, line=TINT_LINE, radius=7)
        text(s, x0 + 8, 336, 123, 40, [{"t": num, "size": 26, "bold": True, "color": GREEN, "font": HEAD, "align": "c"}])
        text(s, x0 + 8, 384, 123, 84, [
            {"t": cap, "size": 12.2, "bold": True, "color": NAVY, "align": "c", "after": 4, "line": 1.0},
            {"t": base, "size": 11.5, "color": BODY, "align": "c", "line": 1.0}])
    text(s, 43.2, 481, 875, 26, [{"t": "Sumber: paparan Sistem Layanan JKN, workshop Healthkathon 20 Okt 2026 (iuran, klaim, "
                                       "badan usaha, cabang); detik.com, 7 Sep 2026 (PKS, 2.429 badan usaha); Perpres 59/2024 "
                                       "Pasal 42. Dampak: benchmark TEPAT pada data sintetis, rata-rata 5 seed.",
                                  "size": 9, "italic": True, "color": GRAY, "line": 1.0}])
    notes(s, N.get(4, ""))

    # 4 Alur proses --------------------------------------------------------
    s = content_slide(prs, "4", "Alur Proses Aplikasi", 4)
    text(s, 140.5, 92, 780, 20, [{"runs": [("Contoh yang diikuti: ", {"bold": True, "color": NAVY}),
                                           ("PT Baru Sejahtera (fiktif), usaha kecil, perdagangan, Kab. Banyusari (fiktif).", {})],
                                  "size": 12, "italic": False, "color": BODY}])
    stages = [
        ("1. Input", "Kepesertaan dan upah per pekerja. Tagihan dan pembayaran iuran. File pertukaran BPJS "
                     "Ketenagakerjaan, 4 kali setahun. Data pelayanan, tanpa diagnosis. Semua sintetis di prototipe."),
        ("2. Proses", "Pencocokan pekerja lintas lembaga: NIK dan kemiripan nama. 4 dari 32 pekerja di BPJS "
                      "Ketenagakerjaan tidak terdaftar di BPJS Kesehatan."),
        ("3. Analisis AI", "Regresi logistik, 29 masukan, dilatih pada 600 pemberi kerja cabang latih. Prioritas: "
                           "peluang pelanggaran dikali perkiraan rupiah, dibagi jam pemeriksaan."),
        ("4. Output", "Peringkat 26 dari 1.500. Potensi Rp46,5 juta (Rp21,8 s.d. 89,7 juta), sintetis. Tiga alasan "
                      "dengan pasal. Saran: pemeriksaan lapangan."),
        ("5. Tindak lanjut", "Pemeriksa memutuskan. Draf surat: LLM, lalu pemeriksa fakta; kembali ke templat bila ada "
                             "fakta baru. Hasil dicatat sebagai calon data latih."),
    ]
    for k, (head, body) in enumerate(stages):
        x0 = 43.6 + k * 179.6
        hb = box(s, x0, 120, 154.8, 46.8, fill=GREEN if k == 2 else NAVY, radius=7)
        text(hb, paras=[{"t": head, "size": 15.9, "bold": True, "color": WHITE, "align": "c"}], anchor="middle")
        if k < 4:
            text(s, x0 + 157, 128, 20, 30, [{"t": ">", "size": 22.3, "bold": True, "color": GREEN, "align": "c"}])
        box(s, x0, 175, 154.8, 178, fill=WHITE, line=PANEL_LINE, radius=7)
        text(s, x0 + 10, 184, 136, 164, [{"t": body, "size": 12.8, "color": BODY, "line": 1.04}])
    bottoms = [
        ("Pengguna / aktor", "Tim pemeriksaan kepatuhan cabang memakai daftar dan kasus. Petugas BPJS Ketenagakerjaan "
                             "ikut saat pemeriksaan bersama."),
        ("Data yang dipakai", "Sintetis, dengan struktur mengikuti data BPJS. Tidak ada data JKN riil. Keterbatasan: "
                              "belum diuji pada data BPJS asli."),
        ("Teknologi", "Python, regresi logistik, aturan berbasis pasal, LLM opsional dengan pemeriksa fakta, Streamlit. "
                      "Regresi logistik: koefisiennya bisa dibaca dan data latih kecil."),
    ]
    for k, (head, body) in enumerate(bottoms):
        x0 = 43.2 + k * 295.2
        label(s, x0, 364, head, w=280)
        box(s, x0, 385, 280.8, 108, fill=WHITE, line=PANEL_LINE, radius=7)
        text(s, x0 + 11, 394, 259, 94, [{"t": body, "size": 13.5, "color": BODY, "line": 1.05}])
    notes(s, N.get(5, ""))

    # 5 Demo -----------------------------------------------------------------
    s = content_slide(prs, "5", "Demo Aplikasi / Prototype", 5)
    box(s, 43.6, 119.2, 547.2, 374.5, fill=NAVY, radius=7)
    s.shapes.add_picture(str(ASSETS / "live_queue.png"), Pt(61.2), Pt(136.8), Pt(511.2), Pt(338.5))
    label(s, 612.3, 114, "Skenario demo")
    steps = [
        ("Daftar prioritas", "anggaran jam, 8% jalur acak, urutan potensi per jam"),
        ("Kasus perusahaan", "PT Baru Sejahtera: tiga alasan dengan pasal, rentang Rp sintetis"),
        ("Mulai tindak lanjut", "tunjuk petugas dan tanggal"),
        ("Draf surat kepatuhan", "LLM opsional dengan pemeriksa fakta; tanpa peluang, peringkat, atau Rp"),
        ("Setujui dan catat hasil", "tidak dikirim sistem; hasil masuk papan Tindak lanjut"),
    ]
    for k, (h1, h2) in enumerate(steps):
        yy = 136 + k * 56
        c = circle(s, 612.4, yy + 9, 30, GREEN)
        text(c, paras=[{"t": str(k + 1), "size": 14, "bold": True, "color": WHITE, "align": "c"}], anchor="middle")
        box(s, 652, yy, 264.6, 49, fill=WHITE, line=PANEL_LINE, radius=7)
        text(s, 661, yy + 3, 250, 44, [{"runs": [(h1 + ": ", {"bold": True, "color": NAVY}), (h2, {})],
                                         "size": 11.8, "color": BODY, "line": 1.0}], anchor="middle")
    label(s, 612.3, 420, "Tautan prototype")
    lb = box(s, 612.0, 440, 304.6, 30, fill=WHITE, line=PANEL_LINE, radius=7)
    text(s, 623, 444, 290, 22, [{"t": "tepat-jkp.streamlit.app", "size": 14, "bold": True, "color": GREEN,
                                 "link": DEMO_URL}], anchor="middle")
    text(s, 612.3, 474, 305, 30, [{"t": "Tangkapan layar aplikasi publik, 28 Okt 2026, cabang uji B, data sintetis.",
                                   "size": 9, "italic": True, "color": GRAY, "line": 1.0}])
    notes(s, N.get(6, ""))

    # 6 Fitur ----------------------------------------------------------------
    s = content_slide(prs, "6", "Penjelasan Fitur", 6)
    feats = [
        ("Daftar prioritas berbasis jam",
         "Mengurutkan pemberi kerja menurut potensi per jam pemeriksaan, sebanyak anggaran jam.",
         "Jam pemeriksa terbatas."),
        ("Kasus, tindak lanjut, dan draf dokumen",
         "Alasan dan pasal per kasus; alur dari mulai sampai hasil; draf surat dengan LLM yang diperiksa fakta.",
         "Pemeriksa butuh bukti yang bisa diperiksa dan surat yang hanya memakai bukti itu."),
        ("Cek kepesertaan saya dan Laporan pekerja",
         "Pekerja mengecek pendaftaran dan iuran, lalu melapor; laporan masuk ke kotak masuk pemeriksa dan bisa "
         "ditautkan ke kasus.",
         "Pekerja bisa baru tahu kepesertaannya nonaktif saat butuh layanan."),
        ("Keadilan & jalur acak, Benchmark",
         "Tingkat terpilih per sektor, skala, dan kab/kota fiktif dengan jumlah dasarnya; 8% jam untuk pilihan acak; "
         "uji lima seed.",
         "Prinsip AI BPJS: uji bias lintas segmen dan wilayah."),
    ]
    for k, (name, fn, prob) in enumerate(feats):
        x0 = 43.6 + (k % 2) * 446.4
        y0 = 112 + (k // 2) * 180
        box(s, x0, y0, 427.0, 172, fill=PANEL, line=PANEL_LINE, radius=7)
        c = box(s, x0 + 18, y0 + 16, 40, 40, fill=GREEN, radius=20)
        text(c, paras=[{"t": str(k + 1), "size": 18, "bold": True, "color": WHITE, "align": "c"}], anchor="middle")
        text(s, x0 + 72, y0 + 16, 340, 40, [{"t": name, "size": 18.5, "bold": True, "color": NAVY}], anchor="middle")
        box(s, x0 + 17.6, y0 + 68, 391.0, 92, fill=WHITE, line=PANEL_LINE, radius=7)
        text(s, x0 + 29, y0 + 75, 370, 80, [
            {"runs": [("Fungsi: ", {"bold": True, "color": NAVY}), (fn, {})], "size": 13.2, "color": BODY,
             "after": 5, "line": 1.03},
            {"runs": [("Menjawab masalah: ", {"bold": True, "color": NAVY}), (prob, {})], "size": 13.2,
             "color": BODY, "line": 1.03}])
    text(s, 43.6, 475, 870, 22, [{"runs": [("Juga ada: ", {"bold": True, "color": NAVY}),
                                           ("Usaha belum terdaftar, daftar terpisah dari data NIB, belum punya alur tindak lanjut.", {})],
                                  "size": 12, "color": BODY}])
    notes(s, N.get(7, ""))

    # + Hasil uji --------------------------------------------------------------
    s = content_slide(prs, "+", "Hasil Uji pada Data Sintetis", 7)
    text(s, 43.2, 106, 870, 22, [{"t": "Cabang uji B (tidak dipakai saat melatih), 5 seed, anggaran jam pemeriksaan yang sama "
                                       "untuk ketiga kebijakan.", "size": 14, "italic": True, "color": GRAY}])
    rows = [
        ("", "TEPAT tanpa jalur acak", "TEPAT + jalur acak 8%", "Aturan pembanding kami"),
        ("Presisi (yang dipilih memang melanggar)", "95,3%", "84,8%", "35,4%"),
        ("Recall (pelanggar yang terpilih)", "33,7%", "32,1%", "8,8%"),
        ("Potensi Rp per jam pemeriksaan", "Rp8,86 jt", "Rp8,63 jt", "Rp3,91 jt"),
        ("Pemberi kerja patuh yang ikut terpilih (dari seluruh yang patuh)", "0,42%", "1,43%", "4,02%"),
    ]
    colw = [250, 140, 140, 140]
    tx, ty = 43.6, 138
    gt = s.shapes.add_table(len(rows), 4, Pt(tx), Pt(ty), Pt(sum(colw)), Pt(46 * len(rows))).table
    tblpr = gt._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblpr.set(attr, "0")
    for j, wv in enumerate(colw):
        gt.columns[j].width = Pt(wv)
    for i, r in enumerate(rows):
        gt.rows[i].height = Pt(46 if i else 40)
        for j, v in enumerate(r):
            cell = gt.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY if i == 0 else (TINT if j == 1 else (PANEL if i % 2 else WHITE)))
            cell.margin_left = cell.margin_right = Pt(8)
            cell.margin_top = cell.margin_bottom = Pt(3)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            text(cell, paras=[{"t": v, "size": 12.5 if i == 0 or j == 0 else 16,
                               "bold": i == 0 or j == 1, "color": WHITE if i == 0 else (GREEN if j == 1 else NAVY if j else BODY),
                               "align": "l" if j == 0 else "c", "line": 1.0}], anchor="middle", inset=6)
    box(s, 735, 138, 181.8, 222, fill=TINT, line=TINT_LINE, radius=7)
    text(s, 745, 152, 162, 50, [{"t": "2,27x", "size": 34, "bold": True, "color": GREEN, "font": HEAD, "align": "c"}])
    text(s, 745, 206, 162, 150, [
        {"t": "potensi Rp per jam dibanding aturan pembanding", "size": 13, "bold": True, "color": NAVY, "align": "c",
         "after": 6, "line": 1.02},
        {"t": "rasio rata-rata lima seed, tanpa jalur acak. Potensi sintetis, bukan uang yang tertagih.",
         "size": 11.5, "color": BODY, "align": "c", "line": 1.02}])
    box(s, 43.6, 382, 871.8, 96, fill=PANEL, line=PANEL_LINE, radius=7)
    label(s, 61.2, 394, "Kondisi gagal yang kami tunjukkan")
    text(s, 61.2, 416, 840, 58, [
        {"t": "Bila pelanggaran jarang (prevalensi 5%, 3 seed), presisi TEPAT turun ke 34,3%, dan pemberi kerja patuh "
              "yang ikut terpilih 5,24% dibanding 3,90% pada aturan. Aturan pembanding adalah aturan yang kami susun "
              "sendiri, bukan cara BPJS memilih sasaran saat ini.", "size": 13.5, "color": BODY, "line": 1.05}])
    text(s, 43.2, 484, 875, 20, [{"t": "Sumber: benchmark TEPAT, results/sweep (lima seed cabang B dan uji sensitivitas), "
                                       "dihitung ulang dari baris mentah 22 Okt 2026.", "size": 9, "italic": True, "color": GRAY}])
    notes(s, "Slide tambahan: hasil uji. Video Script v1 tidak punya narasi khusus untuk slide ini; angkanya sama "
             "dengan narasi Dampak di bagian 4. Dipakai di PDF dan untuk tanya jawab, atau tampil di layar saat narasi "
             "Dampak.\n\nSumber angka: Evidence Bank E78 (lima seed, cabang B) dan E79 (prevalensi 5%, 3 seed).")

    # + Model bisnis dan tata kelola -------------------------------------------
    s = content_slide(prs, "+", "Model Bisnis dan Tata Kelola Data", 8)
    cols = [
        ("Model bisnis (usulan)", [
            "Modul internal BPJS Kesehatan, dijalankan SPPTI (Enterprise AI atau JKN Data Hub). Bukan produk yang dijual.",
            "Pengguna: tim pemeriksaan kepatuhan cabang; petugas BPJS Ketenagakerjaan saat pemeriksaan bersama.",
            "Uji coba bayangan: 1 s.d. 2 cabang, sekitar 3 bulan, data minimum 5 ekstrak.",
            "Biaya: 46 s.d. 91 hari-orang (perkiraan tim, belum angka BPJS).",
            "Diperluas hanya bila lulus gerbang yang disepakati: presisi kasus selesai, temuan per jam pemeriksa, jam "
            "pada pemberi kerja patuh, keadilan antarwilayah, operasional. Pembanding: cara cabang memilih sasaran pada "
            "bulan pertama uji coba.",
        ]),
        ("Tata kelola data", [
            "Prototipe hanya memakai data sintetis. Tidak ada data peserta JKN riil.",
            "Data peserta hanya diolah AI internal BPJS.",
            "DPIA sebelum data asli dipakai (UU 27/2022 tentang Pelindungan Data Pribadi).",
            "Pemeriksa yang memutuskan. TEPAT tidak menjatuhkan sanksi.",
            "Surat ke pemberi kerja tanpa peluang, peringkat, atau rupiah (diuji otomatis).",
        ]),
    ]
    for k, (head, items) in enumerate(cols):
        x0 = 43.6 + k * 446.4
        box(s, x0, 112, 427.0, 384, fill=TINT if k == 0 else PANEL, line=TINT_LINE if k == 0 else PANEL_LINE, radius=7)
        text(s, x0 + 18, 126, 390, 28, [{"t": head, "size": 19, "bold": True, "color": NAVY}])
        text(s, x0 + 18, 162, 392, 326, [{"t": it, "size": 15.5, "color": BODY, "bullet": True, "after": 10, "line": 1.06}
                                         for it in items])
    notes(s, N.get(8, ""))

    # + Batasan ---------------------------------------------------------------
    s = content_slide(prs, "+", "Batasan yang Kami Nyatakan", 9)
    text(s, 43.2, 108, 870, 24, [{"t": "Tiga hal yang belum dibuktikan prototipe ini.",
                                  "size": 15, "italic": True, "color": GRAY}])
    lims = [
        ("Data sintetis", "Semua angka berasal dari data sintetis; belum diuji pada data BPJS asli."),
        ("Belum dikonfirmasi BPJS", "Akses data untuk uji coba dan pemilik operasional belum dikonfirmasi BPJS."),
        ("Prototipe, belum produksi", "Autentikasi, hak akses, dan koneksi ke sistem BPJS belum dibangun."),
    ]
    for k, (h1, body) in enumerate(lims):
        x0 = 43.6 + k * 295.2
        box(s, x0, 150, 280.8, 230, fill=TINT, line=TINT_LINE, radius=7)
        c = box(s, x0 + 18, 168, 40, 40, fill=GREEN, radius=20)
        text(c, paras=[{"t": str(k + 1), "size": 18, "bold": True, "color": WHITE, "align": "c"}], anchor="middle")
        text(s, x0 + 18, 222, 248, 50, [{"t": h1, "size": 18.5, "bold": True, "color": NAVY, "line": 1.0}])
        text(s, x0 + 18, 262, 248, 110, [{"t": body, "size": 17.5, "color": BODY, "line": 1.08}])
    box(s, 43.6, 396, 871.8, 82, fill=PANEL, line=PANEL_LINE, radius=7)
    label(s, 61.2, 408, "Langkah berikut")
    text(s, 61.2, 430, 840, 44, [{"t": "Uji coba bayangan di 1 s.d. 2 cabang, dengan gerbang yang disepakati bersama BPJS, "
                                       "sebelum TEPAT dipakai lebih luas.", "size": 16, "color": BODY, "line": 1.05}])
    notes(s, N.get(9, ""))

    # Terima kasih --------------------------------------------------------------
    s = navy_slide(prs)
    text(s, 64.8, 116, 820, 74, [{"t": "Terima kasih", "size": 51.9, "bold": True, "color": WHITE, "font": HEAD}])
    text(s, 64.8, 196, 820, 30, [{"t": "TEPAT, Tim Johannes Kepler, SMA Negeri 34 Jakarta", "size": 17, "color": PALE}])
    navy_card(s, 65.2, 259.6, 266.4, 108.0, "Tautan prototype / demo", [
        {"t": "tepat-jkp.streamlit.app", "size": 15.9, "bold": True, "color": WHITE, "link": DEMO_URL}])
    navy_card(s, 349.6, 259.6, 266.4, 108.0, "Tautan repositori", [
        {"t": "github.com/rabbaniputrowi-alt/TEPAT", "size": 13.5, "bold": True, "color": WHITE, "link": REPO_URL}])
    navy_card(s, 634.0, 259.6, 266.4, 108.0, "Kontak tim", [
        {"t": "Rabbani Putrowi (ketua)", "size": 15.9, "bold": True, "color": WHITE, "after": 2},
        {"t": "rabbaniputrowi@gmail.com", "size": 13.5, "color": PALE, "link": "mailto:rabbaniputrowi@gmail.com"}])
    text(s, 64.8, 484, 820, 22, [{"t": TAGLINE, "size": 13.7, "italic": True, "color": MINT}])
    notes(s, "Penutup. Musik masuk setelah kata terakhir narasi bagian 9, di kartu ini saja.\n\n"
             "[konfirmasi] Email kontak ketua: rabbaniputrowi@gmail.com.")
    return prs


def all_text(prs) -> list[tuple[int, str]]:
    out = []
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append((i, sh.text_frame.text))
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        out.append((i, c.text_frame.text))
        if slide.has_notes_slide:
            out.append((i, slide.notes_slide.notes_text_frame.text))
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--script", type=Path, default=DEFAULT_SCRIPT)
    ap.add_argument("--out", type=Path, default=OUT)
    a = ap.parse_args()
    prs = build(a.script)
    bad = [(i, BANNED[ch]) for i, t in all_text(prs) for ch in t if ch in BANNED]
    empty_notes = [i for i, s in enumerate(prs.slides, 1) if not s.notes_slide.notes_text_frame.text.strip()]
    a.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(a.out)
    print({"slides": len(prs.slides), "out": str(a.out), "banned_glyphs": bad, "slides_without_notes": empty_notes})
