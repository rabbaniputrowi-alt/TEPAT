"""Evidence-grounded drafting: inspection brief, compliance letter and instalment (cicilan) offer.

Drafts are built from OBSERVED evidence only: the scorer's ranked row (rank, tier, P(pelanggaran), Rp
estimate, action), the feature row (signals) and the employer registry. The synthetic answer key is
never read here.

Generation is hybrid. A deterministic template always works without any key. An optional rewrite by an
OpenAI-compatible chat-completions endpoint only polishes the wording; `check()` then verifies that every
number, date, name and legal citation in the rewrite already appears in the evidence facts or the
template. On any violation, error or timeout the template is returned together with the reasons.

This module never reads environment variables or Streamlit secrets: the caller passes an `LLMConfig`.
"""
from __future__ import annotations

import io
import json
import re
import urllib.error
import urllib.request
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, timedelta

import pandas as pd

from . import scorer

BRIEF, COMPLIANCE, CICILAN = "brief", "surat_kepatuhan", "surat_cicilan"
DOC_LABELS = {
    BRIEF: "Brief pemeriksaan internal",
    COMPLIANCE: "Surat kepatuhan untuk pemberi kerja",
    CICILAN: "Surat penawaran cicilan",
}
LABEL = "DRAF. Data sintetis, tidak dikirim. Disusun oleh prototipe TEPAT untuk demo Healthkathon 2026."
BANNED_GLYPHS = ("\u2014", "\u2013", "\u00b7", "\u2026")
DEFAULT_DEADLINE_DAYS = 14
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober",
         "November", "Desember"]
SKALA = {"micro": "Mikro", "small": "Kecil", "medium": "Menengah", "large": "Besar"}
CONTACT = ("Kontak: [NAMA PETUGAS], [TELEPON ATAU EMAIL KANTOR CABANG]. "
           "Placeholder sintetis, isi sebelum dipakai.")


class DraftNotAvailable(ValueError):
    """The requested document type does not apply to this case."""


# --------------------------------------------------------------------------- evidence

# Each scorer.reasons() text is recognised by a stable phrase, so drafts follow the scorer's order and
# thresholds without copying its weights. tests/test_drafting.py asserts that every reason maps to a kind.
KINDS = [
    ("rws", "sebelum rawat inap"),
    ("care", "biaya perawatan"),
    ("selective", "Membayar BPJS Ketenagakerjaan tepat waktu"),
    ("piggy", "hanya tercatat sebagai tanggungan"),
    ("unreg", "tidak terdaftar di BPJS Kesehatan (setelah pencocokan"),
    ("phk", "dilaporkan PHK"),
    ("clinic", "klinik"),
    ("wage2c", "upah yang dilaporkan ke BPJS TK lebih tinggi"),
    ("wage_peer", "pola sektor sejenis"),
    ("lag", "Median keterlambatan pendaftaran"),
    ("no_joiners", "Tidak ada pendaftaran pekerja baru"),
    ("onoff", "dinonaktifkan di BPJS Kesehatan"),
    ("group_move", "pindah ke segmen mandiri/PBI"),
    ("arrears", "Tunggakan iuran BPJS Kesehatan rata-rata"),
]

QUESTIONS = {
    "rws": "Kapan tanggal mulai kerja pekerja yang didaftarkan menjelang rawat inap, dan mengapa pendaftaran baru "
           "dilakukan saat itu?",
    "care": "Apakah perusahaan mengetahui bahwa kepesertaan pekerja nonaktif saat mereka berobat, dan bagaimana "
            "rencana penyelesaian biaya tersebut?",
    "selective": "Mengapa iuran BPJS Ketenagakerjaan dibayar tepat waktu, sedangkan iuran BPJS Kesehatan terlambat "
                 "atau kurang?",
    "piggy": "Apakah pekerja yang tercatat sebagai tanggungan pasangan memang bekerja di perusahaan ini, dan sejak "
             "kapan?",
    "unreg": "Siapa saja pekerja aktif yang belum didaftarkan ke BPJS Kesehatan, dan apa status hubungan kerjanya "
             "(tetap, kontrak, atau proyek)?",
    "phk": "Apakah pekerja yang dilaporkan PHK benar-benar berhenti bekerja? Minta dokumen PHK dan kontrak lanjutan.",
    "clinic": "Apa hubungan peserta mandiri atau PBI di klinik perusahaan dengan perusahaan, dan apakah sebagian dari "
              "mereka adalah pekerja?",
    "wage2c": "Upah mana yang benar bagi pekerja yang upahnya berbeda di BPJS TK dan BPJS Kesehatan? Cocokkan dengan "
              "slip gaji.",
    "wage_peer": "Apakah upah yang dilaporkan sudah mencakup upah pokok dan tunjangan tetap sesuai slip gaji?",
    "lag": "Bagaimana prosedur pendaftaran pekerja baru, dan mengapa pendaftaran melewati batas 7 hari kerja?",
    "no_joiners": "Apakah ada pekerja baru dalam 12 bulan terakhir, dan mengapa belum ada yang didaftarkan ke BPJS "
                  "Kesehatan?",
    "onoff": "Mengapa pekerja dinonaktifkan dari BPJS Kesehatan sementara masih bekerja?",
    "group_move": "Apakah pekerja yang pindah ke segmen mandiri atau PBI masih bekerja di perusahaan?",
    "arrears": "Apa penyebab tunggakan iuran, dan apakah perusahaan memerlukan program pembayaran bertahap?",
}


def kind_of(reason_text: str) -> str | None:
    for kind, phrase in KINDS:
        if phrase in reason_text:
            return kind
    return None


def _months(share: float) -> int:
    return int(min(12, max(0, round(float(share) * 12))))


def _letter_point(kind: str, f: pd.Series) -> str:
    """Neutral, employer-facing wording of one discrepancy. Counts only: no Rp, no percentages."""
    n = int
    if kind == "rws":
        return (f"{n(f.rws_n)} pekerja didaftarkan ke BPJS Kesehatan tidak lama sebelum menjalani rawat inap, sedangkan "
                "data BPJS Ketenagakerjaan mencatat tanggal mulai kerja yang lebih awal.")
    if kind == "care":
        return (f"Tercatat {n(f.care_susp_n)} kali layanan kesehatan bagi pekerja pada saat kepesertaan mereka nonaktif "
                "karena iuran belum dibayar.")
    if kind == "selective":
        return (f"Iuran BPJS Kesehatan tercatat terlambat atau kurang bayar pada {_months(f.late_ks + f.partial_ks)} "
                "dari 12 bulan terakhir, sementara iuran BPJS Ketenagakerjaan tercatat lancar.")
    if kind == "piggy":
        return (f"{n(f.piggy_n)} pekerja yang tercatat aktif di BPJS Ketenagakerjaan terdaftar di BPJS Kesehatan sebagai "
                "anggota keluarga peserta lain, bukan sebagai pekerja perusahaan.")
    if kind == "unreg":
        return (f"{n(f.tk_unreg_n)} dari {n(f.tk_n)} pekerja yang tercatat di BPJS Ketenagakerjaan belum kami temukan "
                "dalam data kepesertaan BPJS Kesehatan perusahaan.")
    if kind == "phk":
        return (f"{n(f.phk_cycle_n)} pekerja dilaporkan berhenti karena PHK di BPJS Kesehatan, tetapi masih tercatat aktif "
                "di BPJS Ketenagakerjaan atau didaftarkan kembali dalam waktu singkat.")
    if kind == "clinic":
        text = f"{n(f.clinic_tk_n + f.clinic_closed_pbpu_n)} peserta mandiri atau PBI terdaftar di klinik milik perusahaan"
        if f.clinic_tk_n:
            text += f", dan {n(f.clinic_tk_n)} di antaranya tercatat sebagai pekerja di BPJS Ketenagakerjaan"
        return text + "."
    if kind == "wage2c":
        return (f"{n(f.wage2c_n)} pekerja memiliki upah yang dilaporkan ke BPJS Ketenagakerjaan lebih tinggi daripada upah "
                "yang dilaporkan ke BPJS Kesehatan.")
    if kind == "wage_peer":
        return ("Upah yang dilaporkan ke BPJS Kesehatan banyak yang tepat di UMK atau lebih rendah daripada pola usaha "
                "sejenis. Mohon pastikan upah yang dilaporkan mencakup upah pokok dan tunjangan tetap.")
    if kind == "lag":
        return (f"Median jeda antara tanggal mulai kerja dan pendaftaran pekerja baru ke BPJS Kesehatan adalah "
                f"{n(f.lag_med)} hari, sedangkan batas pendaftaran adalah 7 hari kerja.")
    if kind == "no_joiners":
        return ("Tidak ada pendaftaran pekerja baru ke BPJS Kesehatan selama 12 bulan terakhir, sedangkan BPJS "
                "Ketenagakerjaan mencatat adanya pekerja baru.")
    if kind == "onoff":
        return (f"{n(f.onoff_n)} pekerja dinonaktifkan dari BPJS Kesehatan, tetapi masih tercatat bekerja di BPJS "
                "Ketenagakerjaan.")
    if kind == "group_move":
        return (f"{n(f.group_move_n)} pekerja beralih ke segmen peserta mandiri atau PBI, tetapi masih tercatat bekerja "
                "di perusahaan menurut BPJS Ketenagakerjaan.")
    if kind == "arrears":
        return "Tercatat iuran BPJS Kesehatan yang belum dibayar penuh dalam 12 bulan terakhir."
    raise KeyError(kind)


@dataclass(frozen=True)
class Evidence:
    """Everything a draft may state. Built from observed data and scorer output, never from the answer key."""
    employer_id: str
    name: str
    sector: str
    kabkota: str
    skala: str
    rank: int
    tier: str
    p_violation: float
    est_rp: float
    rp_low: float
    rp_high: float
    queue: str
    action: str
    pay_label: str
    hours: int
    reasons: tuple[tuple[str, str, str], ...]          # (kind, inspector text, legal basis), strongest first
    letter_points: tuple[tuple[str, str], ...]         # (neutral employer-facing text, legal basis)
    questions: tuple[str, ...]
    documents: tuple[str, ...]
    ks_late_months: int
    ks_partial_months: int
    ks_arrears_12m_rp: float
    headcount_declining: bool


def evidence(employer_id: str, ranked_row: pd.Series, features_row: pd.Series, registry_row: pd.Series,
             sector_label: str | None = None) -> Evidence:
    f = features_row
    reasons = []
    for _, text, basis in scorer.reasons(f):
        reasons.append((kind_of(text) or "lain", text, basis))
    points = tuple((_letter_point(k, f), basis) for k, _, basis in reasons if k != "lain")
    questions = tuple(dict.fromkeys(QUESTIONS[k] for k, _, _ in reasons if k in QUESTIONS))
    return Evidence(
        employer_id=str(employer_id), name=str(registry_row["name"]),
        sector=str(sector_label or registry_row["sector"]), kabkota=str(registry_row["kabkota"]),
        skala=SKALA.get(str(registry_row.get("skala_usaha", "")), str(registry_row.get("skala_usaha", ""))),
        rank=int(ranked_row["rank"]), tier=str(ranked_row["tier"]), p_violation=float(ranked_row["p_violation"]),
        est_rp=float(ranked_row["est_rp"]), rp_low=float(ranked_row["rp_low"]), rp_high=float(ranked_row["rp_high"]),
        queue=str(ranked_row["queue"]), action=str(ranked_row["action"]), pay_label=str(ranked_row["pay_label"] or ""),
        hours=int(f["hours"]), reasons=tuple(reasons), letter_points=points, questions=questions,
        documents=tuple(scorer.brief(f)), ks_late_months=_months(f["late_ks"]),
        ks_partial_months=_months(f["partial_ks"]), ks_arrears_12m_rp=float(f["rp_arrears"]) * 12,
        headcount_declining=bool(f["decline"] > 0.1))


def available_doc_types(ev: Evidence) -> list[str]:
    """Brief always; compliance letter only when there is a discrepancy to describe; cicilan only in that queue."""
    out = [BRIEF]
    if ev.letter_points and ev.queue != "cicilan":
        out.append(COMPLIANCE)
    if ev.queue == "cicilan":
        out.append(CICILAN)
    return out


# --------------------------------------------------------------------------- templates

def rp(x: float) -> str:
    return "Rp" + f"{x:,.0f}".replace(",", ".")


def tanggal(d: date) -> str:
    return f"{d.day} {BULAN[d.month - 1]} {d.year}"


@dataclass(frozen=True)
class Draft:
    doc_type: str
    title: str
    text: str
    headings: tuple[str, ...]
    facts: tuple[str, ...]         # evidence lines given to the LLM and used by check()


def _assemble(title: str, blocks: list[tuple[str | None, list[str]]]) -> tuple[str, tuple[str, ...]]:
    lines = [LABEL, "", title]
    heads = []
    for head, body in blocks:
        lines.append("")
        if head:
            lines.append(head)
            heads.append(head)
        lines.extend(body)
    return "\n".join(lines).strip() + "\n", tuple(heads)


def _letter_head(ev: Evidence, issued: date, perihal: str) -> list[str]:
    return ["Nomor: [DIISI PETUGAS]", f"Tanggal: {tanggal(issued)}", "", "Kepada Yth.", f"Pimpinan {ev.name}",
            f"Nomor pemberi kerja {ev.employer_id}", ev.kabkota, "", f"Perihal: {perihal}"]


SIGNATURE = ["Hormat kami,", "[NAMA PEJABAT]", "BPJS Kesehatan Kantor Cabang [NAMA KANTOR CABANG]"]


def _brief(ev: Evidence) -> Draft:
    title = "BRIEF PEMERIKSAAN INTERNAL"
    queue = "cicilan" if ev.queue == "cicilan" else "pemeriksaan"
    reasons = []
    for i, (_, text, basis) in enumerate(ev.reasons[:5], 1):
        reasons += [f"{i}. {text}", f"   Dasar: {basis}"]
    action = [f"- {ev.action}"] + ([f"- Catatan pembayaran: {ev.pay_label}"] if ev.pay_label else [])
    blocks = [
        (None, ["Hanya untuk petugas. Jangan dibagikan kepada pemberi kerja."]),
        ("Identitas pemberi kerja", [f"- Nama: {ev.name}", f"- Nomor pemberi kerja: {ev.employer_id}",
                                     f"- Sektor: {ev.sector}", f"- Kab/Kota: {ev.kabkota}",
                                     f"- Skala usaha (OSS): {ev.skala}",
                                     f"- Perkiraan waktu pemeriksaan: {ev.hours} jam"]),
        ("Prioritas menurut model", [f"- Peringkat {ev.rank}, tingkat {ev.tier}, antrean {queue}",
                                     f"- P(pelanggaran): {ev.p_violation:.0%}",
                                     f"- Perkiraan Rp dapat dipulihkan: {rp(ev.est_rp)} "
                                     f"(rentang {rp(ev.rp_low)} s.d. {rp(ev.rp_high)})"]),
        ("Alasan utama", reasons or ["Tidak ada sinyal pelanggaran. Jalur hijau."]),
        ("Pertanyaan saat kunjungan", [f"{i}. {q}" for i, q in enumerate(ev.questions[:6], 1)]
         or ["Tidak ada pertanyaan khusus."]),
        ("Dokumen yang diminta", [f"- {d}" for d in ev.documents]),
        ("Tindakan yang disarankan", action),
        ("Catatan", ["- Disusun dari data yang teramati: sinyal skor, data kepesertaan, dan registri. Kunci jawaban "
                     "sintetis tidak dipakai.",
                     "- Model hanya memberi prioritas. Keputusan pemeriksaan tetap pada petugas."]),
    ]
    text, heads = _assemble(title, blocks)
    facts = [f"Nama pemberi kerja: {ev.name}", f"Nomor pemberi kerja: {ev.employer_id}", f"Sektor: {ev.sector}",
             f"Kab/Kota: {ev.kabkota}", f"Skala usaha: {ev.skala}", f"Waktu pemeriksaan: {ev.hours} jam",
             f"Peringkat: {ev.rank}", f"Tingkat: {ev.tier}", f"P(pelanggaran): {ev.p_violation:.0%}",
             f"Perkiraan Rp: {rp(ev.est_rp)}", f"Rentang: {rp(ev.rp_low)} s.d. {rp(ev.rp_high)}",
             f"Tindakan: {ev.action}"] + ([f"Catatan pembayaran: {ev.pay_label}"] if ev.pay_label else [])
    facts += [f"Alasan: {t} (Dasar: {b})" for _, t, b in ev.reasons[:5]] + [f"Dokumen: {d}" for d in ev.documents]
    return Draft(BRIEF, title, text, heads, tuple(facts))


def _compliance(ev: Evidence, deadline_days: int, issued: date) -> Draft:
    if not ev.letter_points:
        raise DraftNotAvailable("Tidak ada perbedaan data yang dapat dijelaskan dalam surat kepatuhan.")
    title = "SURAT PERMINTAAN KLARIFIKASI DATA KEPESERTAAN"
    due = issued + timedelta(days=deadline_days)
    points = []
    for i, (text, basis) in enumerate(ev.letter_points[:6], 1):
        points += [f"{i}. {text}", f"   Dasar: {basis}"]
    docs = [d for d in ev.documents if d != "Tidak ada permintaan dokumen khusus"] or [
        "Daftar pekerja aktif beserta tanggal mulai kerja"]
    blocks = [
        (None, _letter_head(ev, issued, "Klarifikasi data kepesertaan dan iuran BPJS Kesehatan")),
        (None, ["Dengan hormat,", "",
                "Dalam pencocokan data rutin antara BPJS Kesehatan dan BPJS Ketenagakerjaan, kami menemukan beberapa "
                f"perbedaan data terkait {ev.name}. Perbedaan data belum tentu berarti pelanggaran. Kami memberikan "
                "kesempatan kepada Bapak/Ibu untuk memeriksa dan, bila perlu, memperbaiki data secara mandiri."]),
        ("Perbedaan data yang kami temukan", points),
        ("Kesempatan koreksi mandiri",
         [f"Mohon dalam {deadline_days} hari kalender sejak tanggal surat ini, paling lambat {tanggal(due)}, Bapak/Ibu:",
          "- memeriksa data di atas, lalu mendaftarkan pekerja atau memperbaiki data upah bila memang belum sesuai, atau",
          "- menyampaikan penjelasan beserta dokumen pendukung bila data kami yang keliru."]),
        ("Dokumen yang mohon dikirimkan", [f"- {d}" for d in docs]),
        (None, ["Koreksi yang diselesaikan dalam batas waktu tersebut akan kami catat sebagai koreksi mandiri. Bila "
                "kami tidak menerima tanggapan, kami dapat melanjutkan dengan pemeriksaan sesuai ketentuan.", "",
                CONTACT, "", *SIGNATURE]),
    ]
    text, heads = _assemble(title, blocks)
    facts = [f"Nama pemberi kerja: {ev.name}", f"Nomor pemberi kerja: {ev.employer_id}", f"Kab/Kota: {ev.kabkota}",
             f"Tanggal surat: {tanggal(issued)}", f"Tenggat: {deadline_days} hari kalender, paling lambat {tanggal(due)}"]
    facts += [f"Perbedaan data: {t} (Dasar: {b})" for t, b in ev.letter_points[:6]] + [f"Dokumen: {d}" for d in docs]
    return Draft(COMPLIANCE, title, text, heads, tuple(facts))


def _cicilan(ev: Evidence, deadline_days: int, issued: date) -> Draft:
    if ev.queue != "cicilan":
        raise DraftNotAvailable("Surat penawaran cicilan hanya untuk perusahaan di antrean cicilan.")
    title = "SURAT PENAWARAN PROGRAM PEMBAYARAN BERTAHAP"
    due = issued + timedelta(days=deadline_days)
    facts_lines = []
    if ev.ks_late_months:
        facts_lines.append(f"- Iuran BPJS Kesehatan dibayar terlambat pada {ev.ks_late_months} dari 12 bulan terakhir.")
    if ev.ks_partial_months:
        facts_lines.append(f"- Iuran BPJS Kesehatan dibayar kurang dari tagihan pada {ev.ks_partial_months} dari 12 "
                           "bulan terakhir.")
    if ev.ks_arrears_12m_rp >= 1:
        facts_lines.append(f"- Total kekurangan bayar iuran BPJS Kesehatan dalam 12 bulan terakhir menurut catatan "
                           f"tagihan: {rp(ev.ks_arrears_12m_rp)}.")
    if ev.headcount_declining:
        facts_lines.append("- Jumlah pekerja yang terdaftar menurun dalam setahun terakhir.")
    blocks = [
        (None, _letter_head(ev, issued, "Penawaran program pembayaran bertahap iuran BPJS Kesehatan")),
        (None, ["Dengan hormat,", "",
                "Kami memahami bahwa kondisi usaha dapat memengaruhi kemampuan membayar iuran tepat waktu. Surat ini "
                "bukan surat teguran. Catatan tagihan kami untuk " + ev.name + " menunjukkan hal berikut:"]),
        ("Catatan pembayaran", facts_lines or ["- Tercatat kesulitan pembayaran iuran dalam 12 bulan terakhir."]),
        ("Penawaran program pembayaran bertahap",
         ["Untuk membantu, kami menawarkan program pembayaran bertahap, dengan pola serupa Program REHAB (Rencana "
          "Pembayaran Bertahap) BPJS Kesehatan. Tunggakan dapat dilunasi secara bertahap sesuai "
          "jadwal yang disepakati bersama, sementara iuran bulan berjalan tetap dibayar. Rincian jadwal dan syarat "
          "program akan kami jelaskan dalam pertemuan.",
          "Pelunasan tunggakan penting bagi pekerja: selama kepesertaan nonaktif, biaya pelayanan kesehatan pekerja "
          "menjadi tanggung jawab pemberi kerja (Dasar: Perpres 82/2018 Ps. 42)."]),
        ("Langkah berikutnya",
         [f"Bila Bapak/Ibu berminat, mohon hubungi kami dalam {deadline_days} hari kalender, paling lambat "
          f"{tanggal(due)}, dengan menyiapkan:",
          "- Daftar pekerja aktif saat ini",
          "- Bukti pembayaran iuran BPJS Kesehatan 12 bulan terakhir",
          "- Usulan kemampuan bayar per bulan"]),
        (None, [CONTACT, "", *SIGNATURE]),
    ]
    text, heads = _assemble(title, blocks)
    facts = [f"Nama pemberi kerja: {ev.name}", f"Nomor pemberi kerja: {ev.employer_id}", f"Kab/Kota: {ev.kabkota}",
             f"Tanggal surat: {tanggal(issued)}", f"Tenggat: {deadline_days} hari kalender, paling lambat {tanggal(due)}"]
    facts += [line.lstrip("- ") for line in facts_lines]
    return Draft(CICILAN, title, text, heads, tuple(facts))


def draft(ev: Evidence, doc_type: str, deadline_days: int = DEFAULT_DEADLINE_DAYS, issued: date | None = None) -> Draft:
    """Deterministic template draft. Raises DraftNotAvailable when the type does not apply to the case."""
    issued = issued or date.today()
    if doc_type == BRIEF:
        return _brief(ev)
    if doc_type == COMPLIANCE:
        return _compliance(ev, int(deadline_days), issued)
    if doc_type == CICILAN:
        return _cicilan(ev, int(deadline_days), issued)
    raise DraftNotAvailable(f"Jenis dokumen tidak dikenal: {doc_type}")


# --------------------------------------------------------------------------- fact checker

NUM = re.compile(r"\d+(?:[.,]\d+)*")
DATE = re.compile(r"\b\d{1,2}\s+(?:" + "|".join(BULAN) + r")\s+\d{4}\b")
COMPANY = re.compile(r"\b(?:PT|CV|UD|PD|Koperasi|Yayasan|Firma)\.?\s+[A-Z][\w&'-]*(?:\s+[A-Z][\w&'-]*)*")
CAPITAL = re.compile(r"\b[A-Z][A-Za-z'-]+\b")
REG_SHORT = re.compile(r"\b(UU|PP|Perpres|Permenkes|Permenaker|Perpu|Perppu|Perbpjs|Peraturan BPJS)\s*(?:No\.?\s*|Nomor\s+)?"
                       r"(\d+)\s*/\s*(\d{4})", re.I)
REG_LONG = re.compile(r"\b(Undang-Undang|Peraturan Pemerintah|Peraturan Presiden|Peraturan Menteri Kesehatan|"
                      r"Peraturan Menteri Ketenagakerjaan|Peraturan BPJS Kesehatan|Peraturan BPJS)\s+(?:Nomor\s+|No\.?\s*)?"
                      r"(\d+)\s+Tahun\s+(\d{4})", re.I)
REG_NAMES = {"undang-undang": "uu", "peraturan pemerintah": "pp", "peraturan presiden": "perpres",
             "peraturan menteri kesehatan": "permenkes", "peraturan menteri ketenagakerjaan": "permenaker",
             "peraturan bpjs kesehatan": "perbpjs", "peraturan bpjs": "perbpjs", "perppu": "perpu"}
ARTICLE_ONE = r"\d+[A-Za-z]?(?:\s*(?:\(\d+\)|ayat\s*\(?\d+\)?))?(?:\s*(?:huruf\s+)?[a-z]\b)?"
ARTICLE = re.compile(r"\b(?:Pasal|Ps\.?)\s*(" + ARTICLE_ONE + r"(?:\s*(?:,|&|dan|jo\.?)\s*" + ARTICLE_ONE + r")*)", re.I)
# Capitalised words a faithful rewrite may use without them appearing in the template.
COMMON = {
    "Anda", "Bapak", "Ibu", "Saudara", "Kami", "Yth", "Dengan", "Hormat", "Mohon", "Bila", "Apabila", "Jika", "Untuk",
    "Atas", "Terima", "Kasih", "Demikian", "Sehubungan", "Selain", "Oleh", "Karena", "Namun", "Sementara", "Selama",
    "Dalam", "Pada", "Dari", "Kepada", "Perihal", "Nomor", "Tanggal", "Lampiran", "Dasar", "Hal", "Data", "Kantor",
    "Cabang", "Badan", "Penyelenggara", "Jaminan", "Sosial", "Kesehatan", "Ketenagakerjaan", "Nasional", "Program",
    "Pasal", "Ayat", "Tahun", "Undang-Undang", "Peraturan", "Pemerintah", "Presiden", "Republik", "Indonesia",
    "Silakan", "Berikut", "Adapun", "Kemudian", "Selanjutnya", "Agar", "Sebagai", "Setelah", "Sebelum", "Kontak",
    "Catatan", "Ini", "Itu", "Hanya", "Tidak", "Semua", "Seluruh", "Perusahaan", "Pekerja", "Iuran", "Surat",
    "BPJS", "JKN", "KIS", "NIK", "PHK", "UMK", "PBI", "TK", "OSS", "DRAF", "TEPAT",
}


def _canon_num(token: str) -> str:
    return token.replace(".", "").replace(",", "")


def numbers(text: str) -> set[str]:
    return {_canon_num(m) for m in NUM.findall(text)}


def _reg_key(kind: str, num: str, year: str) -> str:
    k = kind.lower().strip()
    k = REG_NAMES.get(k, k).replace(" ", "")
    return f"{k} {int(num)}/{year}"


def _articles(group: str) -> set[str]:
    out = set()
    for one in re.findall(ARTICLE_ONE, group):
        base = re.match(r"\d+[A-Za-z]?", one).group(0).lower()   # "13" or "13a"
        out.add(f"ps {base}")
        ayat = re.search(r"\((\d+)\)|ayat\s*\(?(\d+)", one)
        if ayat:
            out.add(f"ps {base}({ayat.group(1) or ayat.group(2)})")
    return out


def citations(text: str) -> set[str]:
    """Normalised regulation and article citations: 'perpres 82/2018', 'ps 42', 'ps 42(2)'."""
    out = {_reg_key(*m) for m in REG_SHORT.findall(text)} | {_reg_key(*m) for m in REG_LONG.findall(text)}
    for m in ARTICLE.finditer(text):
        out |= _articles(m.group(1))
    return out


def _words(text: str) -> set[str]:
    return set(re.findall(r"[A-Za-z][A-Za-z'-]*", text))


@dataclass
class Check:
    ok: bool
    problems: list[str] = field(default_factory=list)


def check(text: str, d: Draft, ev: Evidence) -> Check:
    """Every number, date, name and legal citation in `text` must already be in the template or the facts."""
    allowed_text = d.text + "\n" + "\n".join(d.facts)
    problems: list[str] = []
    if not text or len(text.strip()) < 0.3 * len(d.text.strip()):
        problems.append("teks kosong atau terlalu pendek")
    for g in BANNED_GLYPHS:
        if g in text:
            problems.append(f"glyph terlarang: U+{ord(g):04X}")
    if LABEL not in text:
        problems.append("label DRAF hilang")
    bad = sorted(numbers(text) - numbers(allowed_text))
    if bad:
        problems.append("angka tidak ada di bukti: " + ", ".join(bad))
    bad = sorted(set(DATE.findall(text)) - set(DATE.findall(allowed_text)))
    if bad:
        problems.append("tanggal tidak ada di bukti: " + ", ".join(bad))
    known_names = {re.sub(r"\s+", " ", m) for m in COMPANY.findall(allowed_text)} | {ev.name}
    bad = sorted({re.sub(r"\s+", " ", m) for m in COMPANY.findall(text)
                  if not any(re.sub(r"\s+", " ", m).startswith(k) for k in known_names)})
    if bad:
        problems.append("nama perusahaan tidak dikenal: " + ", ".join(bad))
    caps = {w for w in _words(allowed_text) if w[0].isupper()}
    lower = {w.lower() for w in _words(allowed_text)}
    bad = sorted({w for w in CAPITAL.findall(text) if w not in caps and w not in COMMON and w.lower() not in lower})
    if bad:
        problems.append("nama atau istilah baru: " + ", ".join(bad))
    bad = sorted(citations(text) - citations(allowed_text))
    if bad:
        problems.append("dasar hukum tidak ada di template: " + ", ".join(bad))
    if d.doc_type in (COMPLIANCE, CICILAN):
        low = text.lower()
        leaks = [w for w in ("peringkat", "probabilitas", "p(pelanggaran)", "model", "skor") if w in low]
        if d.doc_type == COMPLIANCE:
            leaks += ["Rp"] if "Rp" in text else []
            leaks += sorted(set(re.findall(r"\d+(?:,\d+)?\s?%", text)) - set(re.findall(r"\d+(?:,\d+)?\s?%", d.text)))
        if leaks:
            problems.append("data internal di surat untuk pemberi kerja: " + ", ".join(leaks))
    return Check(not problems, problems)


# --------------------------------------------------------------------------- optional LLM rewrite

@dataclass(frozen=True)
class LLMConfig:
    api_key: str = field(repr=False)
    base_url: str
    model: str
    timeout_s: float = 20.0

    @staticmethod
    def from_values(values: dict) -> LLMConfig | None:
        """Build from TEPAT_LLM_API_KEY, TEPAT_LLM_BASE_URL, TEPAT_LLM_MODEL; None unless all three are set."""
        key, url, model = (str(values.get(k) or "").strip()
                           for k in ("TEPAT_LLM_API_KEY", "TEPAT_LLM_BASE_URL", "TEPAT_LLM_MODEL"))
        return LLMConfig(key, url, model) if key and url and model else None


def chat_completion(config: LLMConfig, messages: list[dict]) -> str:
    """POST {base_url}/chat/completions (OpenAI-compatible). Plain urllib, bounded by config.timeout_s."""
    body = json.dumps({"model": config.model, "messages": messages, "temperature": 0.2}).encode("utf-8")
    req = urllib.request.Request(config.base_url.rstrip("/") + "/chat/completions", data=body, method="POST",
                                 headers={"Content-Type": "application/json",
                                          "Authorization": f"Bearer {config.api_key}"})
    with urllib.request.urlopen(req, timeout=config.timeout_s) as resp:
        payload = json.loads(resp.read().decode("utf-8"))
    return str(payload["choices"][0]["message"]["content"])


def prompt(d: Draft) -> list[dict]:
    system = ("Anda penyunting surat dinas BPJS Kesehatan. Perbaiki kejelasan dan kesantunan bahasa saja. Jangan "
              "menambah, menghapus, atau mengubah fakta, angka, nama, tanggal, maupun dasar hukum. Jangan menambah nama "
              "orang atau perusahaan. Pertahankan baris pertama (label DRAF) persis sama, pertahankan judul, judul "
              "bagian, dan daftar berbutir. Jangan memakai tanda pisah panjang, tanda pisah sedang, titik tengah, atau "
              "elipsis. Keluarkan hanya teks dokumen.")
    user = ("FAKTA BUKTI (satu-satunya sumber fakta):\n" + "\n".join(f"- {x}" for x in d.facts)
            + "\n\nDRAF TEMPLATE:\n" + d.text)
    return [{"role": "system", "content": system}, {"role": "user", "content": user}]


def _unfence(text: str) -> str:
    """Drop a surrounding Markdown code fence if the model added one."""
    t = text.strip()
    if t.startswith("```"):
        t = t.split(chr(10), 1)[1] if chr(10) in t else ""
        if t.rstrip().endswith("```"):
            t = t.rstrip()[:-3]
    return t.strip() + chr(10)


@dataclass
class Result:
    text: str
    source: str                    # "template" or "llm"
    problems: list[str] = field(default_factory=list)
    note: str = ""


def rewrite(d: Draft, ev: Evidence, config: LLMConfig | None = None,
            client: Callable[[list[dict]], str] | None = None) -> Result:
    """Optional LLM polish of a template draft. Falls back to the template on any problem, with the reasons."""
    if client is None and config is None:
        return Result(d.text, "template", [], "LLM tidak dikonfigurasi. Template deterministik dipakai.")
    call = client or (lambda m: chat_completion(config, m))
    try:
        out = call(prompt(d))
    except (TimeoutError, OSError, urllib.error.URLError) as e:   # socket timeouts are OSError subclasses
        kind = "batas waktu habis" if "timed out" in str(e).lower() or isinstance(e, TimeoutError) else type(e).__name__
        return Result(d.text, "template", [f"panggilan LLM gagal: {kind}"], "Kembali ke template.")
    except Exception as e:  # malformed response or anything else: never surface payloads or keys
        return Result(d.text, "template", [f"panggilan LLM gagal: {type(e).__name__}"], "Kembali ke template.")
    out = out.strip().strip("`").strip() + "\n"
    verdict = check(out, d, ev)
    if not verdict.ok:
        return Result(d.text, "template", verdict.problems, "Hasil LLM ditolak pemeriksa fakta. Kembali ke template.")
    return Result(out, "llm", [], "Bahasa diperhalus LLM dan lolos pemeriksaan fakta.")


# --------------------------------------------------------------------------- DOCX

def to_docx(text: str, title: str, headings: tuple[str, ...] = ()) -> bytes:
    """Render a draft's plain text to DOCX bytes. Title and section headings become Word headings."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if s == LABEL:
            doc.add_paragraph().add_run(s).bold = True
        elif s == title:
            doc.add_heading(s, level=1)
        elif s in headings:
            doc.add_heading(s, level=2)
        elif s.startswith("- "):
            doc.add_paragraph(s[2:], style="List Bullet")
        elif line.startswith("   "):
            doc.add_paragraph(s).paragraph_format.left_indent = Pt(18)
        else:
            doc.add_paragraph(s)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
